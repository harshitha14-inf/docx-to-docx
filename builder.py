from docx import Document

from models import (
    Paragraph,
    Table,
    Image,
    Caption
)


class Builder:

    def build(self, model, out_file):

        doc = Document()

        # -------------------
        # Headers
        # -------------------

        if model.headers:

            section = doc.sections[0]

            header = section.header

            header.paragraphs[0].text = (
                model.headers[0].text
            )

        # -------------------
        # Footers
        # -------------------

        if model.footers:

            section = doc.sections[0]

            footer = section.footer

            footer.paragraphs[0].text = (
                model.footers[0].text
            )

        # -------------------
        # Main Content
        # -------------------

        for item in model.content:

            # Paragraph

            if isinstance(item, Paragraph):

                doc.add_paragraph(
                    item.text
                )

            # Caption

            elif isinstance(item, Caption):

                p = doc.add_paragraph()

                run = p.add_run(
                    item.text
                )

                run.bold = True

            # Table

            elif isinstance(item, Table):

                if not item.data:
                    continue

                rows = len(item.data)

                cols = max(
                    len(row)
                    for row in item.data
                )

                table = doc.add_table(
                    rows=rows,
                    cols=cols
                )

                for r, row in enumerate(
                    item.data
                ):

                    for c, value in enumerate(
                        row
                    ):

                        table.cell(
                            r,
                            c
                        ).text = value

            # Image

            elif isinstance(item, Image):

                try:

                    doc.add_picture(
                        item.path
                    )

                except Exception as e:

                    print(
                        f"Could not add image "
                        f"{item.name}: {e}"
                    )

        doc.save(out_file)

