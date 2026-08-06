from collections import Counter
import hashlib
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from lxml import etree

from .style_engine.style_mapper import StyleMapper


CAPTION_RE = re.compile(
    r"^(figure|fig\.)\s+\d+(?:\.\d+)*(?::|\b)",
    re.IGNORECASE,
)

TABLE_CAPTION_RE = re.compile(
    r"^table\s+\d+(?:\.\d+)*(?::|\b)",
    re.IGNORECASE,
)

# Mirrors extractor.py's formula detection - a paragraph carrying a heading
# style/outline level whose text is an equation must never count as a
# heading in either the source or output snapshot.
FORMULA_FUNCTION_RE = re.compile(
    r"\b(ATAN2|ATAN|SIN|COS|TAN|SQRT|LOG)\s*\(",
    re.IGNORECASE,
)

FORMULA_ASSIGNMENT_RE = re.compile(
    r"[A-Za-z_][A-Za-z0-9_]*(?:<[^>\s]{1,20}>)?\s*(?:\+=|-=|\*=|/=|=)\s*[\w(<\-.]"
)


def _is_formula_text(text):

    if not text:
        return False

    if FORMULA_FUNCTION_RE.search(text):
        return True

    # Anchored at the start - see extractor.py's _is_formula_text for why.
    if FORMULA_ASSIGNMENT_RE.match(text.strip()):
        return True

    return False

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def validate(
    source_docx,
    output_docx,
    template_used=None,
    style_mapper=None,
    formatting_policy_path="config/formatting_policy.json",
    acceptance_rules_path="config/acceptance_rules.json",
    strict_roundtrip=True,
):

    source_snapshot = _snapshot(source_docx)
    output_snapshot = _snapshot(output_docx)

    checks = [
        "image_count",
        "image_format_distribution",
        "image_relationship_count",
        "section_count",
        "landscape_section_count",
        "orientation_sequence",
        "figure_count",
        "caption_count",
        "heading_count",
        "caption_image_adjacency",
        "content_order_integrity",
    ]

    report = {}
    failed = 0

    print("\n========== VALIDATION REPORT ==========\n")
    print("CRITICAL CHECKS")
    print("----------------")

    for item in checks:

        source_value = source_snapshot[item]
        output_value = output_snapshot[item]

        passed = _critical_check_pass(
            item,
            source_value,
            output_value,
            strict_roundtrip,
        )

        report[item] = {
            "source": source_value,
            "output": output_value,
            "pass": passed,
            "severity": "CRITICAL" if strict_roundtrip else "INFO",
        }

        if passed:
            print(
                f"[PASS] {item:<28} "
                f"Source={source_value} "
                f"Output={output_value}"
            )
        else:
            failed += 1
            print(
                f"[FAIL] {item:<28} "
                f"Source={source_value} "
                f"Output={output_value}"
            )

    overall_status = "PASS" if failed == 0 else "FAIL"

    print("\n======================================")
    print(f"Overall Status : {overall_status}")
    print(f"Critical Failed: {failed}")
    print("======================================\n")

    mapper = style_mapper or StyleMapper.from_file(Path("config/style_map.json"))
    formatting_policy = _load_json_or_default(Path(formatting_policy_path), {})
    acceptance_rules = _load_json_or_default(Path(acceptance_rules_path), {})

    compliance = _build_contract_compliance(
        source_snapshot,
        output_snapshot,
        output_docx,
        template_used,
        mapper,
        formatting_policy,
        acceptance_rules,
    )

    report["template_compliance"] = compliance

    print("CONTRACT COMPLIANCE")
    print("-------------------")
    print(
        f"[{'PASS' if compliance['rules']['loss_rules_pass'] else 'FAIL'}] Loss Rules"
    )
    print(
        f"[{'PASS' if compliance['rules']['heading_rules_pass'] else 'FAIL'}] Heading Rules"
    )
    print(
        f"[{'PASS' if compliance['rules']['caption_rules_pass'] else 'FAIL'}] Caption Rules"
    )
    print(
        f"[{'PASS' if compliance['rules']['template_rules_pass'] else 'FAIL'}] Template Rules"
    )
    print(
        f"[{'PASS' if compliance['rules']['formula_safety_pass'] else 'FAIL'}] Formula Safety"
    )
    print(
        f"[{'PASS' if compliance['forbidden_content_hits'] == 0 else 'FAIL'}] Forbidden Content"
    )

    report["summary"] = {
        "status": overall_status,
        "critical_failed": failed,
        "infineon_compliance": "PASS"
        if overall_status == "PASS"
        and compliance["overall"] == "PASS"
        else "FAIL",
    }

    return report


def _snapshot(docx_path):

    document = Document(docx_path)

    with ZipFile(docx_path) as archive:

        root = etree.fromstring(archive.read("word/document.xml"))

        rid_to_target = _extract_relationships(archive)

        header_footer_count = _header_footer_non_empty_count_from_archive(archive)

    body = root.xpath("./w:body", namespaces=NS)[0]

    top_level_items = []
    order_index = 0

    for element in body:

        if etree.QName(element).localname == "sectPr":
            continue

        top_level_items.append(
            _classify_top_level_item(
                element,
                rid_to_target,
                order_index,
            )
        )
        order_index += 1

    grouped_items = _group_figure_blocks(top_level_items)

    image_refs = _collect_all_image_refs(body, rid_to_target)

    format_distribution = Counter(image_ref["image_format"] for image_ref in image_refs)

    orientation_sequence = []

    for section in document.sections:
        if section.orientation == WD_ORIENT.LANDSCAPE:
            orientation_sequence.append("landscape")
        else:
            orientation_sequence.append("portrait")

    content_signature = [_content_signature(item) for item in grouped_items]

    hyperlinks_external, hyperlinks_internal, cross_refs, bookmarks = _extract_reference_counts(root)

    return {
        "image_count": len(image_refs),
        "image_format_distribution": dict(sorted(format_distribution.items())),
        "image_relationship_count": len({image_ref["relationship_id"] for image_ref in image_refs}),
        "section_count": len(document.sections),
        "landscape_section_count": sum(1 for orientation in orientation_sequence if orientation == "landscape"),
        "orientation_sequence": orientation_sequence,
        "figure_count": sum(1 for item in grouped_items if item["type"] == "figure"),
        "caption_count": sum(1 for item in top_level_items if item["type"] == "caption"),
        "heading_count": sum(1 for item in top_level_items if item["type"] == "heading"),
        "formula_count": sum(1 for item in top_level_items if item["type"] == "formula"),
        "suspicious_heading_count": sum(
            1
            for item in top_level_items
            if item["type"] == "heading" and _is_formula_text(item.get("text", ""))
        ),
        "table_count": sum(1 for item in top_level_items if item["type"] == "table"),
        "textbox_count": sum(1 for item in top_level_items if item["type"] == "textbox"),
        "reference_count": hyperlinks_external + hyperlinks_internal + cross_refs + bookmarks,
        "header_footer_count": header_footer_count,
        "caption_image_adjacency": sum(1 for item in grouped_items if item["type"] == "figure"),
        "heading_styles": [
            item.get("style")
            for item in top_level_items
            if item["type"] == "heading" and item.get("style")
        ],
        "caption_style_items": [
            {
                "caption_type": item.get("caption_type", "image"),
                "style": item.get("style"),
            }
            for item in top_level_items
            if item["type"] == "caption"
        ],
        "content_order_integrity": hashlib.sha1(
            json.dumps(content_signature, separators=(",", ":")).encode("utf-8")
        ).hexdigest(),
    }


def _extract_relationships(archive):

    rid_to_target = {}
    rels_path = "word/_rels/document.xml.rels"

    if rels_path not in archive.namelist():
        return rid_to_target

    rels_root = ET.fromstring(archive.read(rels_path))

    for rel in rels_root:

        r_id = rel.get("Id")
        target = rel.get("Target", "")

        if r_id and target:
            rid_to_target[r_id] = target

    return rid_to_target


def _classify_top_level_item(element, rid_to_target, order_index):

    local_name = etree.QName(element).localname

    if local_name == "tbl":
        return {
            "type": "table",
            "order_index": order_index,
        }

    if local_name != "p":
        return {
            "type": local_name,
            "order_index": order_index,
        }

    image_refs = _collect_image_refs_from_element(element, rid_to_target, order_index)

    if image_refs:
        return {
            "type": "image",
            "order_index": order_index,
            "image_refs": image_refs,
        }

    textbox_texts = _get_textbox_texts(element)
    text = _get_visible_paragraph_text(element).strip()
    style_name = _get_paragraph_style_name(element)

    if CAPTION_RE.match(text):
        return {
            "type": "caption",
            "caption_type": "image",
            "order_index": order_index,
            "text": text,
            "style": style_name,
        }

    if TABLE_CAPTION_RE.match(text):
        return {
            "type": "caption",
            "caption_type": "table",
            "order_index": order_index,
            "text": text,
            "style": style_name,
        }

    if _is_formula_text(text) or (textbox_texts and not text and any(_is_formula_text(t) for t in textbox_texts)):
        return {
            "type": "formula",
            "order_index": order_index,
            "text": text or "\n".join(textbox_texts),
            "style": style_name,
        }

    heading_level = _detect_heading_level(element)

    if heading_level is not None:
        return {
            "type": "heading",
            "order_index": order_index,
            "level": heading_level,
            "text": text,
            "style": style_name,
        }

    if textbox_texts and not text:
        return {
            "type": "textbox",
            "order_index": order_index,
            "textbox_count": len(textbox_texts),
        }

    return {
        "type": "paragraph",
        "order_index": order_index,
        "text": text,
        "style": style_name,
    }


def _group_figure_blocks(items):

    grouped_items = []
    index = 0

    while index < len(items):

        item = items[index]

        if item["type"] == "caption" and item.get("caption_type") == "image":

            images = []
            look_ahead = index + 1

            while look_ahead < len(items) and items[look_ahead]["type"] == "image":
                images.append(items[look_ahead])
                look_ahead += 1

            if images:
                grouped_items.append(
                    {
                        "type": "figure",
                        "order_index": item["order_index"],
                        "caption_position": "before",
                        "caption": item,
                        "images": images,
                    }
                )
                index = look_ahead
                continue

        if item["type"] == "image":

            images = [item]
            look_ahead = index + 1

            while look_ahead < len(items) and items[look_ahead]["type"] == "image":
                images.append(items[look_ahead])
                look_ahead += 1

            if (
                look_ahead < len(items)
                and items[look_ahead]["type"] == "caption"
                and items[look_ahead].get("caption_type") == "image"
            ):
                grouped_items.append(
                    {
                        "type": "figure",
                        "order_index": item["order_index"],
                        "caption_position": "after",
                        "caption": items[look_ahead],
                        "images": images,
                    }
                )
                index = look_ahead + 1
                continue

        grouped_items.append(item)
        index += 1

    return grouped_items


def _collect_all_image_refs(body, rid_to_target):

    image_refs = []

    for r_id in body.xpath(".//*[local-name()='blip']/@r:embed", namespaces=NS):
        image_refs.append(_build_image_ref(r_id, rid_to_target, -1))

    for r_id in body.xpath(".//*[local-name()='imagedata']/@r:id", namespaces=NS):
        image_refs.append(_build_image_ref(r_id, rid_to_target, -1))

    return image_refs


def _collect_image_refs_from_element(element, rid_to_target, order_index):

    image_refs = []
    seen = set()

    for r_id in element.xpath(".//*[local-name()='blip']/@r:embed", namespaces=NS):

        if r_id in seen:
            continue

        seen.add(r_id)
        image_refs.append(_build_image_ref(r_id, rid_to_target, order_index))

    for r_id in element.xpath(".//*[local-name()='imagedata']/@r:id", namespaces=NS):

        if r_id in seen:
            continue

        seen.add(r_id)
        image_refs.append(_build_image_ref(r_id, rid_to_target, order_index))

    return image_refs


def _build_image_ref(r_id, rid_to_target, order_index):

    target = rid_to_target.get(r_id, "")

    return {
        "name": Path(target).name,
        "relationship_id": r_id,
        "image_format": Path(target).suffix.lower().lstrip("."),
        "document_position_index": order_index,
        "target": target,
    }


def _get_textbox_texts(element):

    textbox_texts = []
    altcontent_seen = set()

    for node in element.xpath(".//*[local-name()='txbxContent']"):

        altcontent = node.xpath("ancestor::*[local-name()='AlternateContent'][1]")

        if altcontent:

            alt_key = id(altcontent[0])

            if alt_key in altcontent_seen:
                continue

            altcontent_seen.add(alt_key)

        text = "".join(node.xpath(".//*[local-name()='t']/text()")).strip()

        if text:
            textbox_texts.append(text)

    return textbox_texts


def _get_visible_paragraph_text(element):

    text_parts = []

    for t_node in element.xpath(".//*[local-name()='t']"):

        if not t_node.text:
            continue

        ancestor = t_node.getparent()
        skip_node = False

        while ancestor is not None and ancestor is not element:

            local_name = etree.QName(ancestor).localname

            if local_name in {"drawing", "txbxContent"}:
                skip_node = True
                break

            ancestor = ancestor.getparent()

        if not skip_node:
            text_parts.append(t_node.text)

    return "".join(text_parts)


def _content_signature(item):

    if item["type"] == "figure":
        image_count = sum(len(image["image_refs"]) for image in item["images"])

        return {
            "type": "figure",
            "caption_position": item["caption_position"],
            "image_count": image_count,
        }

    if item["type"] == "image":
        return {
            "type": "image",
            "image_count": len(item["image_refs"]),
        }

    if item["type"] == "textbox":
        return {
            "type": "textbox",
            "textbox_count": item.get("textbox_count", 0),
        }

    return {
        "type": item["type"],
    }


def _get_paragraph_style_name(element):

    style_value = element.xpath("./*[local-name()='pPr']/*[local-name()='pStyle']/@*[local-name()='val']")

    if not style_value:
        return None

    return style_value[0]


def _detect_heading_level(element):

    style_name = (_get_paragraph_style_name(element) or "").strip().lower()

    match = re.search(r"heading\s*([1-9])$", style_name)

    if match:
        return int(match.group(1))

    outline_levels = element.xpath(
        "./*[local-name()='pPr']/*[local-name()='outlineLvl']/@*[local-name()='val']"
    )

    if outline_levels:
        try:
            return int(outline_levels[0]) + 1
        except ValueError:
            return None

    return None


def _extract_reference_counts(root):

    hyperlinks_external = len(root.xpath(".//w:hyperlink[@r:id]", namespaces=NS))
    hyperlinks_internal = len(root.xpath(".//w:hyperlink[@w:anchor]", namespaces=NS))

    instr_texts = [
        node.text.strip()
        for node in root.xpath(".//w:instrText", namespaces=NS)
        if node.text
    ]

    hyperlinks_external += sum(
        1 for text in instr_texts
        if re.search(r"\bHYPERLINK\b", text) and r"\l" not in text
    )

    hyperlinks_internal += sum(
        1 for text in instr_texts
        if re.search(r"\bHYPERLINK\b", text) and r"\l" in text
    )

    cross_refs = sum(
        1 for text in instr_texts
        if re.search(r"\b(REF|PAGEREF)\b", text)
    )

    bookmarks = len(
        [
            bookmark
            for bookmark in root.xpath(".//w:bookmarkStart", namespaces=NS)
            if not bookmark.get(f"{{{NS['w']}}}name", "").startswith("_")
        ]
    )

    return hyperlinks_external, hyperlinks_internal, cross_refs, bookmarks


def _header_footer_non_empty_count(document):

    count = 0

    for section in document.sections:
        header_has_text = any(p.text.strip() for p in section.header.paragraphs)
        footer_has_text = any(p.text.strip() for p in section.footer.paragraphs)

        if header_has_text:
            count += 1

        if footer_has_text:
            count += 1

    return count


def _header_footer_non_empty_count_from_archive(archive):

    count = 0

    for filename in archive.namelist():

        if not filename.startswith("word/header") and not filename.startswith("word/footer"):
            continue

        if not filename.endswith(".xml"):
            continue

        try:
            root = etree.fromstring(archive.read(filename))
        except Exception:
            continue

        text = "".join(root.xpath(".//*[local-name()='t']/text()"))

        if text.strip():
            count += 1

    return count


def _load_json_or_default(path, default_value):

    if not path.exists():
        return default_value

    with open(path, "r", encoding="utf-8") as handle:
        return json.load(handle)


def _critical_check_pass(item, source_value, output_value, strict_roundtrip):

    if strict_roundtrip:
        return source_value == output_value

    if item == "image_format_distribution":
        for fmt, source_count in source_value.items():
            if output_value.get(fmt, 0) < source_count:
                return False
        return True

    if item in {
        "image_count",
        "image_relationship_count",
        "section_count",
        "landscape_section_count",
        "figure_count",
        "caption_count",
        "heading_count",
        "caption_image_adjacency",
    }:
        return output_value >= source_value

    if item in {"orientation_sequence", "content_order_integrity"}:
        return True

    return source_value == output_value


def _build_contract_compliance(
    source_snapshot,
    output_snapshot,
    output_docx,
    template_used,
    mapper,
    formatting_policy,
    acceptance_rules,
):

    output_document = Document(output_docx)
    style_lookup = _build_style_lookup(output_document)

    style_names = {style.name for style in output_document.styles if style.name}

    required_styles = set(formatting_policy.get("required_styles", []))
    optional_styles = set(formatting_policy.get("optional_styles", []))
    required_style_missing = sorted(required_styles - style_names)
    optional_style_missing = sorted(optional_styles - style_names)

    forbidden_content = formatting_policy.get("forbidden_content", [])
    forbidden_content_hits = _count_forbidden_content_hits(output_document, forbidden_content)

    heading_styles = [
        _normalize_style_name(style, style_lookup)
        for style in output_snapshot.get("heading_styles", [])
    ]
    expected_heading_styles = {
        mapper.style_for_heading(level)
        for level in range(1, 7)
        if mapper.style_for_heading(level)
    }

    heading_match_count = sum(1 for style in heading_styles if style in expected_heading_styles)
    heading_total = len(heading_styles)
    heading_style_match_pct = _percentage(heading_match_count, heading_total)

    caption_style_items = output_snapshot.get("caption_style_items", [])
    figure_caption_expected = mapper.style_for_caption("image")
    table_caption_expected = mapper.style_for_caption("table")

    figure_caption_items = [item for item in caption_style_items if item["caption_type"] == "image"]
    table_caption_items = [item for item in caption_style_items if item["caption_type"] == "table"]

    figure_caption_match_pct = _percentage(
        sum(
            1
            for item in figure_caption_items
            if _normalize_style_name(item.get("style"), style_lookup) == figure_caption_expected
        ),
        len(figure_caption_items),
    )

    table_caption_match_pct = _percentage(
        sum(
            1
            for item in table_caption_items
            if _normalize_style_name(item.get("style"), style_lookup) == table_caption_expected
        ),
        len(table_caption_items),
    )

    losses = {
        "image_loss": max(0, source_snapshot["image_count"] - output_snapshot["image_count"]),
        "table_loss": max(0, source_snapshot["table_count"] - output_snapshot["table_count"]),
        "heading_loss": max(0, source_snapshot["heading_count"] - output_snapshot["heading_count"]),
        "reference_loss": max(0, source_snapshot["reference_count"] - output_snapshot["reference_count"]),
        "textbox_loss": max(0, source_snapshot["textbox_count"] - output_snapshot["textbox_count"]),
        "header_footer_loss": max(0, source_snapshot["header_footer_count"] - output_snapshot["header_footer_count"]),
        "landscape_section_loss": max(
            0,
            source_snapshot["landscape_section_count"] - output_snapshot["landscape_section_count"],
        ),
    }

    template_load_failures = 0

    if not template_used or not Path(template_used).exists():
        template_load_failures = 1

    rules = acceptance_rules
    loss_rules = rules.get("migration_loss", {})
    caption_rules = rules.get("caption_rules", {})
    heading_rules = rules.get("heading_rules", {})
    template_rules = rules.get("template_rules", {})

    loss_rules_pass = all(
        losses.get(key, 0) <= value
        for key, value in loss_rules.items()
    )

    heading_rules_pass = heading_style_match_pct >= heading_rules.get("heading_style_match", 0)

    caption_rules_pass = (
        figure_caption_match_pct >= caption_rules.get("figure_caption_style_match", 0)
        and table_caption_match_pct >= caption_rules.get("table_caption_style_match", 0)
    )

    template_rules_pass = (
        len(required_style_missing) <= template_rules.get("required_style_missing", 0)
        and template_load_failures <= template_rules.get("template_load_failures", 0)
    )

    # Hard fail - formulas must never survive as headings (they must never be
    # numbered, never enter the TOC, never participate in outline/navigation).
    formula_safety_pass = output_snapshot.get("suspicious_heading_count", 0) == 0

    overall = "PASS"

    if (
        not loss_rules_pass
        or not heading_rules_pass
        or not caption_rules_pass
        or not template_rules_pass
        or not formula_safety_pass
        or forbidden_content_hits > 0
    ):
        overall = "FAIL"

    return {
        "template_used": template_used,
        "losses": losses,
        "style_match": {
            "heading_style_match": heading_style_match_pct,
            "figure_caption_style_match": figure_caption_match_pct,
            "table_caption_style_match": table_caption_match_pct,
        },
        "required_style_missing": required_style_missing,
        "optional_style_missing": optional_style_missing,
        "template_load_failures": template_load_failures,
        "forbidden_content_hits": forbidden_content_hits,
        "suspicious_heading_count": output_snapshot.get("suspicious_heading_count", 0),
        "rules": {
            "loss_rules_pass": loss_rules_pass,
            "heading_rules_pass": heading_rules_pass,
            "caption_rules_pass": caption_rules_pass,
            "template_rules_pass": template_rules_pass,
            "formula_safety_pass": formula_safety_pass,
        },
        "overall": overall,
    }


def _percentage(matched, total):

    if total == 0:
        return 100

    return round((matched / total) * 100, 2)


def _count_forbidden_content_hits(document, forbidden_content):

    if not forbidden_content:
        return 0

    hits = 0
    needles = [value.lower() for value in forbidden_content]

    for paragraph in document.paragraphs:
        text = paragraph.text.strip().lower()

        if not text:
            continue

        if any(needle in text for needle in needles):
            hits += 1

    return hits


def _build_style_lookup(document):

    lookup = {}

    for style in document.styles:
        if style.name:
            lookup[style.name.strip().lower()] = style.name

        style_id = getattr(style, "style_id", None)
        if style_id:
            lookup[str(style_id).strip().lower()] = style.name or str(style_id)

    return lookup


def _normalize_style_name(style_value, style_lookup):

    if not style_value:
        return style_value

    normalized = style_lookup.get(str(style_value).strip().lower())

    if normalized:
        return normalized

    return style_value


if __name__ == "__main__":

    import sys

    if len(sys.argv) != 3:

        print(
            "Usage:\n"
            "python validator.py source.docx output.docx"
        )

        sys.exit(1)

    result = validate(sys.argv[1], sys.argv[2])

    print(json.dumps(result, indent=4))
