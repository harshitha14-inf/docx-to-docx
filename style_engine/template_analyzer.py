import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE


def analyze_template(template_path, output_path="templates/style_inventory.json"):
    template = Path(template_path)

    if not template.exists():
        raise FileNotFoundError(f"Template not found: {template}")

    document = Document(template)

    inventory = {
        "template": str(template),
        "styles": sorted(style.name for style in document.styles if style.name),
        "paragraph_styles": sorted(
            style.name
            for style in document.styles
            if style.name and style.type == WD_STYLE_TYPE.PARAGRAPH
        ),
        "character_styles": sorted(
            style.name
            for style in document.styles
            if style.name and style.type == WD_STYLE_TYPE.CHARACTER
        ),
        "table_styles": sorted(
            style.name
            for style in document.styles
            if style.name and style.type == WD_STYLE_TYPE.TABLE
        ),
        "header_footer": _extract_header_footer_info(document),
    }

    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    with open(out_file, "w", encoding="utf-8") as handle:
        json.dump(inventory, handle, indent=2)

    return inventory


def _extract_header_footer_info(document):
    section_data = []

    for idx, section in enumerate(document.sections):
        header_styles = sorted(
            {
                paragraph.style.name
                for paragraph in section.header.paragraphs
                if paragraph.style is not None and paragraph.style.name
            }
        )
        footer_styles = sorted(
            {
                paragraph.style.name
                for paragraph in section.footer.paragraphs
                if paragraph.style is not None and paragraph.style.name
            }
        )

        section_data.append(
            {
                "section_index": idx,
                "header_text_present": any(p.text.strip() for p in section.header.paragraphs),
                "footer_text_present": any(p.text.strip() for p in section.footer.paragraphs),
                "header_styles": header_styles,
                "footer_styles": footer_styles,
            }
        )

    return section_data


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("template_path")
    parser.add_argument("--output", default="templates/style_inventory.json")

    args = parser.parse_args()

    result = analyze_template(args.template_path, args.output)
    print(json.dumps(result, indent=2))
